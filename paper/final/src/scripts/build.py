"""论文 Word 生成：各章 Markdown → 数字与文献填充 → pandoc → 按《格式规范》后处理。

    python paper/final/src/scripts/build.py [out.docx]

默认输出 paper/final/竞赛论文.docx。数字只来自 src/data/facts.json（由 facts.py 从
results/ 提取），任何未替换的 {{KEY}} 或未知文献键都会让构建失败。
"""

from __future__ import annotations

import copy
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
SRC = HERE.parent
FINAL = SRC.parent
MD, FIG, DATA = SRC / 'md', SRC / 'figs', SRC / 'data'
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else FINAL / '竞赛论文.docx'
WORK = Path(tempfile.gettempdir()) / 'paper_final_build'
WORK.mkdir(parents=True, exist_ok=True)
PANDOC = 'pandoc'

TITLE = '基于跨核同步层次分层的通用 NPU 多核切图与调度'
SONG, HEI, KAI, TNR = '宋体', '黑体', '楷体', 'Times New Roman'
XIAOSI, SIHAO, SANHAO, XIAOER, XIAOYI = 12, 14, 16, 18, 24
PAGE_W, MARGIN = 21.0, 2.5
TEXT_W = PAGE_W - 2 * MARGIN


# --------------------------------------------------------------------------
# 1. 参考样式文档
# --------------------------------------------------------------------------

def set_fonts(obj, east=SONG, ascii_=TNR, size=None, bold=None):
    f = obj.font
    f.name = ascii_
    rpr = obj.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for k in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rfonts.set(qn(k), ascii_)
    rfonts.set(qn('w:eastAsia'), east)
    for k in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(k)) is not None:
            del rfonts.attrib[qn(k)]
    if size:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    f.color.rgb = RGBColor(0, 0, 0)
    f.italic = False


def pstyle(doc, name, base='Normal', east=SONG, size=XIAOSI, bold=False,
           align=None, indent_first=None, before=0, after=0, ascii_=TNR):
    try:
        st = doc.styles[name]
    except KeyError:
        st = doc.styles.add_style(name, 1)
        st.base_style = doc.styles[base]
    set_fonts(st, east=east, size=size, bold=bold, ascii_=ascii_)
    pf = st.paragraph_format
    pf.line_spacing = 1.0
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if align is not None:
        pf.alignment = align
    pf.first_line_indent = indent_first
    return st


def para_border(st, sides):
    ppr = st.element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr')
    for side, sz in sides.items():
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), str(sz))
        e.set(qn('w:space'), '1')
        e.set(qn('w:color'), '000000')
        bdr.append(e)
    ppr.append(bdr)


def make_reference():
    ref = WORK / 'ref.docx'
    with ref.open('wb') as fh:
        subprocess.run([PANDOC, '--print-default-data-file', 'reference.docx'],
                       stdout=fh, check=True)
    doc = docx.Document(str(ref))
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, side, Cm(MARGIN))
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.5)
    J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    body = dict(align=J, indent_first=Pt(24))
    pstyle(doc, 'Normal', size=XIAOSI)
    for n in ('Body Text', 'First Paragraph'):
        pstyle(doc, n, **body)
    pstyle(doc, 'Compact', size=XIAOSI, indent_first=Pt(0))
    # 一级标题：四号黑体居中；其余标题用小四号宋体加粗（其他汉字一律小四号宋体）
    h1 = pstyle(doc, 'Heading 1', east=HEI, size=SIHAO, align=C, before=6, after=12)
    h1.paragraph_format.page_break_before = True
    h1.paragraph_format.keep_with_next = True
    h2 = pstyle(doc, 'Heading 2', size=XIAOSI, bold=True, align=L, before=9, after=6)
    h2.paragraph_format.keep_with_next = True
    h3 = pstyle(doc, 'Heading 3', size=XIAOSI, bold=True, align=L, before=6, after=3)
    h3.paragraph_format.keep_with_next = True
    for n in ('Heading 1', 'Heading 2', 'Heading 3'):
        st = doc.styles[n]
        st.paragraph_format.first_line_indent = Pt(0)
        st.paragraph_format.left_indent = Pt(0)
    pstyle(doc, 'Image Caption', size=XIAOSI, align=C, indent_first=Pt(0), before=3, after=9)
    tcap = pstyle(doc, 'Table Caption', size=XIAOSI, align=C, indent_first=Pt(0),
                  before=9, after=3)
    tcap.paragraph_format.keep_with_next = True
    for n in ('Captioned Figure', 'Figure'):
        st = pstyle(doc, n, size=XIAOSI, align=C, indent_first=Pt(0), before=6)
        st.paragraph_format.keep_with_next = True
    eq = pstyle(doc, 'Equation', size=XIAOSI, align=L, indent_first=Pt(0), before=3, after=3)
    ts = eq.paragraph_format.tab_stops
    ts.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    ts.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    for n, size, after in (('CoverOrg', XIAOER, 6), ('CoverTitle', XIAOYI, 6)):
        pstyle(doc, n, east=KAI, size=size, bold=True, align=C, indent_first=Pt(0),
               before=6, after=after)
    # 摘要页抬头：与封面同字体，字号略小，保证摘要不超过两页
    pstyle(doc, 'AbsOrg', east=KAI, size=SIHAO, bold=True, align=C, indent_first=Pt(0),
           before=0, after=3)
    pstyle(doc, 'AbsTitle', east=KAI, size=SANHAO, bold=True, align=C, indent_first=Pt(0),
           before=0, after=0)
    pstyle(doc, 'AbstractHead', east=HEI, size=SIHAO, align=C, indent_first=Pt(0),
           before=6, after=6)
    pstyle(doc, 'Keywords', size=XIAOSI, align=J, indent_first=Pt(0), before=6)
    rf = pstyle(doc, 'Reference', size=XIAOSI, align=L, indent_first=Pt(0), after=3)
    rf.paragraph_format.left_indent = Pt(24)
    rf.paragraph_format.first_line_indent = Pt(-24)
    at = pstyle(doc, 'AlgoTitle', size=XIAOSI, align=L, indent_first=Pt(0), before=6, after=2)
    para_border(at, {'top': 12, 'bottom': 6})
    at.paragraph_format.keep_with_next = True
    ast = pstyle(doc, 'AlgoStep', size=XIAOSI, align=J, indent_first=Pt(0))
    ast.paragraph_format.left_indent = Pt(18)
    ast.paragraph_format.first_line_indent = Pt(-18)
    ast.paragraph_format.keep_with_next = True
    ae = pstyle(doc, 'AlgoEnd', size=XIAOSI, align=J, indent_first=Pt(0), after=6)
    ae.paragraph_format.left_indent = Pt(18)
    ae.paragraph_format.first_line_indent = Pt(-18)
    para_border(ae, {'bottom': 12})
    pstyle(doc, 'Source Code', east=SONG, ascii_='Consolas', size=9, align=L,
           indent_first=Pt(0))
    try:
        set_fonts(doc.styles['Verbatim Char'], east=SONG, ascii_='Consolas', size=10.5)
    except KeyError:
        pass
    bq = pstyle(doc, 'Block Text', size=XIAOSI, align=J, indent_first=Pt(0), before=3, after=3)
    bq.paragraph_format.left_indent = Pt(21)
    bq.paragraph_format.right_indent = Pt(21)
    pstyle(doc, 'Footer', size=10.5, align=C, indent_first=Pt(0))
    try:
        doc.styles['Hyperlink'].font.color.rgb = RGBColor(0, 0, 0)
    except KeyError:
        pass
    doc.save(str(ref))
    return ref


# --------------------------------------------------------------------------
# 2. Markdown 预处理：数字、文献、公式、标记
# --------------------------------------------------------------------------

TAB = '`<w:r><w:tab/></w:r>`{=openxml}'
CITE = re.compile(r'\[(@[\w\-]+(?:\s*;\s*@[\w\-]+)*)\]')


def marker(name):
    return f'\n::: {{custom-style="Marker"}}\n@@{name}@@\n:::\n'


def md_escape(s):
    return re.sub(r'([\\`*_\[\]<>#$])', r'\\\1', s)


def fill_facts(text):
    facts = json.loads((DATA / 'facts.json').read_text(encoding='utf-8'))
    used, missing = [], []

    def sub(m):
        k = m.group(1)
        if k not in facts:
            missing.append(k)
            return m.group(0)
        if k not in used:
            used.append(k)
        # 不用千位分隔逗号（按国家标准 GB/T 15835 的数字书写习惯）
        return re.sub(r'(?<=\d),(?=\d{3}(?!\d))', '', str(facts[k]['v']))

    text = re.sub(r'\{\{([A-Za-z0-9_.]+)\}\}', sub, text)
    if missing:
        sys.exit(f'facts.json 中缺少：{sorted(set(missing))}')
    (DATA / 'used_facts.json').write_text(json.dumps(used, ensure_ascii=False, indent=0),
                                          encoding='utf-8')
    return text


def fill_citations(text):
    refs = json.loads((DATA / 'refs.json').read_text(encoding='utf-8'))
    order = []

    def sub(m):
        keys = re.findall(r'@([\w\-]+)', m.group(1))
        out = []
        for k in keys:
            if k not in refs:
                sys.exit(f'refs.json 中没有文献键：{k}')
            if k not in order:
                order.append(k)
            out.append(f'\\[{order.index(k) + 1}\\]')
        return ''.join(out)

    text = CITE.sub(sub, text)
    unused = [k for k in refs if k not in order]
    if unused:
        sys.exit(f'refs.json 中有未被正文引用的文献：{unused}')
    lst = '\n'.join(f'::: {{custom-style="Reference"}}\n\\[{i + 1}\\] {md_escape(refs[k]["text"])}\n:::\n'
                    for i, k in enumerate(order))
    text = text.replace('@@REFERENCES@@', lst)
    (DATA / 'ref_order.json').write_text(json.dumps(order, ensure_ascii=False), encoding='utf-8')
    return text


def cn_quotes(text):
    """正文中的 ASCII 双引号成对换成中文引号（跳过代码、公式与属性块）。"""
    out, fence, bad = [], False, []
    for line in text.split('\n'):
        if line.lstrip().startswith('```'):
            fence = not fence
        if fence or line.lstrip().startswith('```'):
            out.append(line)
            continue
        masks = []

        def mask(mo):
            masks.append(mo.group(0))
            return f'\x00{len(masks) - 1}\x00'

        s = re.sub(r'`[^`]*`|\$\$?[^$]*\$\$?|\{[^}]*\}', mask, line)
        s = re.sub(r'"([^"\n]*)"', r'“\1”', s)
        if '"' in s:
            bad.append(line[:40])
        out.append(re.sub(r'\x00(\d+)\x00', lambda mo: masks[int(mo.group(1))], s))
    if bad:
        sys.exit(f'引号不成对：{bad}')
    return '\n'.join(out)


def math_bars(text):
    """公式兼容处理：绝对值竖线统一写成 \\left|…\\right|，Word 与 LibreOffice 都能正确显示。"""
    def fix(m):
        s = m.group(0)
        s = re.sub(r'(?<!\\left)(?<!\\right)(?<!\\bigl)(?<!\\bigr)\|([^|$]+?)\|', r'\\left|\1\\right|', s)
        # 另两处写法 Word 能显示、LibreOffice 预览会报错，换成等价写法
        for a, b in (('\\lvert', '\\left|'), ('\\rvert', '\\right|'),
                     ('\\bigl|', '\\left|'), ('\\bigr|', '\\right|'),
                     ('\\ast', '\\text{*}'), ('\\setminus', '\\backslash ')):
            s = s.replace(a, b)
        return re.sub(r'(?<!\\)\\ ', r'\\;', s)  # 控制空格在预览中显示为方框

    parts = re.split(r'(^```.*?^```)', text, flags=re.M | re.S)
    return ''.join(p if p.startswith('```') else
                   re.sub(r'\$\$.+?\$\$|\$[^$\n]+?\$', fix, p, flags=re.S) for p in parts)


def preprocess():
    parts = sorted(MD.glob('*.md'))
    text = '\n\n'.join(p.read_text(encoding='utf-8') for p in parts)
    text = fill_facts(text)
    text = fill_citations(text)
    text = cn_quotes(text)
    text = math_bars(text)
    text = text.replace('@@FIG@@', str(FIG))

    def eq(m):
        body, num = m.group(1).strip(), m.group(2)
        return (f'::: {{custom-style="Equation"}}\n{TAB}$\\displaystyle {body}$'
                f'{TAB}（{num}）\n:::')

    text = re.sub(r'^\$\$(.+?)\$\$\s*\{#([\d\-]+)\}\s*$', eq, text, flags=re.M | re.S)

    # 表格单元格内的数值括注改用半角括号，便于窄列排版
    def tabline(m):
        return re.sub(r'(\d)（([−+\-\d][^）]*)）', r'\1 (\2)', m.group(0))

    text = re.sub(r'^\|.*\|\s*$', tabline, text, flags=re.M)
    for name in ('COVER', 'TITLELINE', 'PAGEBREAK', 'SECTIONBREAK'):
        text = text.replace(f'@@{name}@@', marker(name))
    left = sorted(set(re.findall(r'@@[A-Z]+@@|\{\{[A-Za-z0-9_.]+\}\}|\[@[\w\-]+', text)))
    left = [x for x in left if x not in {f'@@{n}@@' for n in
                                         ('COVER', 'TITLELINE', 'PAGEBREAK', 'SECTIONBREAK')}]
    if left:
        sys.exit(f'仍有未替换的标记：{left}')
    (WORK / 'paper.md').write_text(text, encoding='utf-8')
    return WORK / 'paper.md'


# --------------------------------------------------------------------------
# 3. 后处理
# --------------------------------------------------------------------------

def run(par, text, east=SONG, size=XIAOSI, bold=False, underline=False, ascii_=TNR):
    r = par.add_run(text)
    set_fonts(r, east=east, size=size, bold=bold, ascii_=ascii_)
    r.font.underline = underline
    return r


def set_cell_border(cell, **kw):
    tcpr = cell._tc.get_or_add_tcPr()
    b = tcpr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders')
        tcpr.append(b)
    for side, sz in kw.items():
        e = OxmlElement(f'w:{side}')
        if sz:
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(sz))
            e.set(qn('w:color'), '000000')
        else:
            e.set(qn('w:val'), 'nil')
        b.append(e)


def three_line(tbl, size=Pt(10.5)):
    tblpr = tbl._tbl.tblPr
    for tag in ('w:tblBorders', 'w:tblStyle', 'w:tblLook'):
        e = tblpr.find(qn(tag))
        if e is not None:
            tblpr.remove(e)
    borders = OxmlElement('w:tblBorders')
    for side, sz in (('top', 12), ('bottom', 12), ('left', 0), ('right', 0),
                     ('insideH', 0), ('insideV', 0)):
        e = OxmlElement(f'w:{side}')
        if sz:
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), str(sz))
            e.set(qn('w:color'), '000000')
        else:
            e.set(qn('w:val'), 'nil')
        borders.append(e)
    tblpr.append(borders)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    w = tblpr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW')
        tblpr.append(w)
    w.set(qn('w:type'), 'pct')
    w.set(qn('w:w'), '5000')
    rows = tbl.rows
    if rows:
        trpr = rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:tblHeader'))
        for c in rows[0].cells:
            set_cell_border(c, bottom=6)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.bold = True
    for row in rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = size


def fld_run(par, kind, instr=None):
    r = OxmlElement('w:r')
    fc = OxmlElement('w:fldChar')
    fc.set(qn('w:fldCharType'), kind)
    r.append(fc)
    par._p.append(r)
    if instr:
        r2 = OxmlElement('w:r')
        it = OxmlElement('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = instr
        r2.append(it)
        r.addnext(r2)


def remove(par):
    par._p.getparent().remove(par._p)


def page_break_para(anchor):
    p = anchor.insert_paragraph_before()
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p._p.append(r)
    remove(anchor)


def section_break(anchor):
    """封面单独成节：无页脚、不编页码。"""
    p = anchor.insert_paragraph_before()
    ppr = p._p.get_or_add_pPr()
    sp = copy.deepcopy(DOC.element.body.find(qn('w:sectPr')))
    for e in list(sp):
        if e.tag in (qn('w:footerReference'), qn('w:headerReference'), qn('w:pgNumType'),
                     qn('w:titlePg')):
            sp.remove(e)
    t = OxmlElement('w:type')
    t.set(qn('w:val'), 'nextPage')
    sp.insert(0, t)
    ppr.append(sp)
    remove(anchor)


def build_cover(anchor):
    lp = anchor.insert_paragraph_before()
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lp.paragraph_format.first_line_indent = Pt(0)
    lp.paragraph_format.space_before = Pt(36)
    # 四枚赛事标志取自《格式规范》文档（承办单位标志只取左侧圆形徽标）
    for i, path in enumerate(sorted((FIG / 'cover').glob('logo*.png'))):
        if i:
            lp.add_run('　')
        lp.add_run().add_picture(str(path), height=Cm(2.2))
    for _ in range(3):
        anchor.insert_paragraph_before().style = DOC.styles['CoverOrg']
    for text, st in (('中国研究生创新实践系列大赛', 'CoverOrg'),
                     ('“华为杯”第二十三届中国研究生', 'CoverTitle'),
                     ('数学建模竞赛', 'CoverTitle')):
        p = anchor.insert_paragraph_before()
        p.style = DOC.styles[st]
        run(p, text, east=KAI, size=XIAOER if st == 'CoverOrg' else XIAOYI, bold=True)
    for _ in range(4):
        anchor.insert_paragraph_before().style = DOC.styles['CoverOrg']
    rows = [['学　　校', ''], ['参赛队号', ''], ['队员姓名', '1.'], ['', '2.'], ['', '3.']]
    widths = [3.5, 8.5]
    tbl = DOC.add_table(rows=len(rows), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.width = Cm(widths[j])
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run(p, val, east=HEI if j == 0 else SONG, size=SIHAO)
            set_cell_border(cell, top=0, left=0, right=0, bottom=6 if j == 1 or i < 2 else 0)
    # 队员姓名三行合并左列
    tbl.cell(2, 0).merge(tbl.cell(4, 0))
    set_cell_border(tbl.cell(2, 0), top=0, left=0, right=0, bottom=6)
    tblpr = tbl._tbl.tblPr
    for tag in ('w:tblStyle',):
        e = tblpr.find(qn(tag))
        if e is not None:
            tblpr.remove(e)
    anchor._p.addprevious(tbl._tbl)
    remove(anchor)


def title_line(anchor):
    p = anchor.insert_paragraph_before()
    p.style = DOC.styles['Normal']
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, '题　目：', east=HEI, size=SANHAO)
    run(p, '　' + TITLE + '　', east=HEI, size=SANHAO, underline=True)
    remove(anchor)


def add_footer_page_numbers():
    sec = DOC.sections[-1]
    sec.footer.is_linked_to_previous = False
    ft = sec.footer
    p = ft.paragraphs[0] if ft.paragraphs else ft.add_paragraph()
    p.style = DOC.styles['Footer']
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld_run(p, 'begin', ' PAGE ')
    fld_run(p, 'separate')
    run(p, '1', size=10.5)
    fld_run(p, 'end')
    sectpr = sec._sectPr
    pg = sectpr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType')
        sectpr.append(pg)
    pg.set(qn('w:start'), '1')
    pg.set(qn('w:fmt'), 'decimal')
    for s in DOC.sections:
        s.different_first_page_header_footer = False
        # 不设页眉：两节的页眉均为空
        for para in s.header.paragraphs:
            for r in list(para.runs):
                r._r.getparent().remove(r._r)


def main():
    global DOC
    ref = make_reference()
    md = preprocess()
    raw = WORK / 'raw.docx'
    subprocess.run([PANDOC, str(md), '-f', 'markdown-smart+tex_math_dollars+raw_attribute',
                    '-t', 'docx', '--reference-doc', str(ref), '-o', str(raw)], check=True)
    DOC = docx.Document(str(raw))
    markers = {}
    for p in DOC.paragraphs:
        t = p.text.strip()
        if t.startswith('@@') and t.endswith('@@'):
            markers.setdefault(t, []).append(p)
    for p in markers.get('@@COVER@@', []):
        build_cover(p)
    for p in markers.get('@@TITLELINE@@', []):
        title_line(p)
    for p in markers.get('@@SECTIONBREAK@@', []):
        section_break(p)
    for p in markers.get('@@PAGEBREAK@@', []):
        page_break_para(p)
    # 表格：封面信息表之外一律三线表；附录 A 大表用 9 磅字
    for tbl in DOC.tables[1:]:
        prev = tbl._tbl.getprevious()
        cap = ''.join(prev.itertext()) if prev is not None else ''
        three_line(tbl, Pt(9) if '表 A' in cap else Pt(10.5))
    for p in DOC.paragraphs:
        if p._p.xpath('.//w:drawing'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            if p.style.name != 'Normal':
                p.paragraph_format.keep_with_next = True
        for r in p.runs:
            if r.style is not None and r.style.name == 'Verbatim Char':
                set_fonts(r, east=SONG, ascii_='Consolas', size=10.5)
    add_footer_page_numbers()
    DOC.core_properties.title = TITLE
    DOC.core_properties.author = ''
    DOC.core_properties.last_modified_by = ''
    DOC.core_properties.comments = ''
    DOC.save(str(OUT))
    left = [p.text for p in DOC.paragraphs if '@@' in p.text or '{{' in p.text]
    if left:
        sys.exit(f'成品中仍有标记：{left[:5]}')
    print('saved', OUT)


if __name__ == '__main__':
    main()
