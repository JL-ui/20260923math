# 本程序及代码是在人工智能工具辅助下完成的。
# 人工智能工具：Claude Sonnet 5（型号 claude-sonnet-5）、Claude Opus 5.5（型号 claude-opus-5-5）
# 开发机构：Anthropic 公司
# 版本发布日期：Claude Sonnet 5 为 2026-06-30，Claude Opus 5.5 为 2026-09-22
"""把论文 Markdown 转成 Word (.docx)。

    python solution/md2docx.py paper/论文.md -o paper/论文.docx

支持：多级标题、正文段落、无序/有序列表、表格、图片（自动居中 + 图注）、
块级公式 `$$...$$`（用 matplotlib mathtext 渲染成图片嵌入）、行内公式
（转成 Unicode 近似文本）、粗体/斜体/行内代码。
渲染失败的公式回退为等宽文本，不会中断转换。
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt                                 # noqa: E402
from docx import Document                                       # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT                  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                   # noqa: E402
from docx.oxml.ns import qn                                     # noqa: E402
from docx.shared import Inches, Pt, RGBColor                    # noqa: E402

MATH_DIR = None

# mathtext 不支持的宏 -> 支持的等价写法
_MATH_FIX = [
    (r'\\Bigl', r'\\left'), (r'\\Bigr', r'\\right'),
    (r'\\bigl', r'\\left'), (r'\\bigr', r'\\right'),
    (r'\\Big\\|', r'\\right|'), (r'\\biggl', r'\\left'), (r'\\biggr', r'\\right'),
    (r'\\mathbb\s*\{?1\}?', r'\\mathbf{1}'),
    (r'\\rightsquigarrow', r'\\to'),
    (r'\\varnothing', r'\\emptyset'),
    (r'\\;', r'\\ '), (r'\\!', r''), (r'\\,', r'\\ '),
    (r'\\displaystyle', r''),
    (r'\\begin\{cases\}', r'\\{'), (r'\\end\{cases\}', r''),
    (r'\\\\', r'\\quad '),
    (r'&', r''),
    (r'\\lesssim', r'\\leq'),
    (r'\\quad\\text\{', r'\\quad \\mathrm{'),
    (r'\\text\{', r'\\mathrm{'),
]

_UNICODE = {
    r'\alpha': 'α', r'\beta': 'β', r'\gamma': 'γ', r'\delta': 'δ',
    r'\lambda': 'λ', r'\mu': 'μ', r'\sigma': 'σ', r'\theta': 'θ',
    r'\rho': 'ρ', r'\pi': 'π', r'\Delta': 'Δ', r'\Phi': 'Φ', r'\kappa': 'κ',
    r'\tau': 'τ', r'\times': '×', r'\le': '≤', r'\ge': '≥', r'\neq': '≠',
    r'\to': '→', r'\in': '∈', r'\cup': '∪', r'\cap': '∩', r'\subseteq': '⊆',
    r'\sum': 'Σ', r'\max': 'max', r'\min': 'min', r'\cdot': '·',
    r'\varnothing': '∅', r'\emptyset': '∅', r'\approx': '≈', r'\forall': '∀',
    r'\exists': '∃', r'\ldots': '…', r'\dots': '…', r'\mathrm': '',
    r'\mathbb': '', r'\text': '', r'\left': '', r'\right': '', r'\;': ' ',
    r'\,': ' ', r'\ ': ' ', r'\bigl': '', r'\bigr': '', r'\Bigl': '',
    r'\Bigr': '', r'\lceil': '⌈', r'\rceil': '⌉', r'\lfloor': '⌊',
    r'\rfloor': '⌋', r'\infty': '∞', r'\lesssim': '≲', r'\gtrsim': '≳',
    r'\mathcal': '', r'\sigma': 'σ', r'\Sigma': 'Σ',
}


def _sanitize_math(tex: str) -> str:
    out = tex.strip()
    for pat, rep in _MATH_FIX:
        out = re.sub(pat, rep, out)
    return out


_CJK = re.compile(r'[　-鿿＀-￯]')


def render_math(tex: str, out_dir: Path, fontsize=15):
    if _CJK.search(tex):
        return None            # mathtext 不支持中日韩字形，退回纯文本
    out_dir.mkdir(parents=True, exist_ok=True)
    key = hashlib.blake2b(tex.encode('utf-8'), digest_size=8).hexdigest()
    path = out_dir / f'eq_{key}.png'
    if path.is_file():
        return path
    body = _sanitize_math(tex)
    for attempt in (body, body.replace('\\quad', ' ')):
        try:
            fig = plt.figure(figsize=(0.01, 0.01))
            fig.text(0, 0, '${}$'.format(attempt), fontsize=fontsize)
            fig.savefig(path, bbox_inches='tight', pad_inches=0.08, dpi=220,
                        facecolor='white')
            plt.close(fig)
            return path
        except Exception:                                       # noqa: BLE001
            plt.close('all')
    return None


def inline_to_text(tex: str) -> str:
    s = tex
    for k, v in sorted(_UNICODE.items(), key=lambda kv: -len(kv[0])):
        s = s.replace(k, v)
    s = re.sub(r'[{}]', '', s)
    s = re.sub(r'\\[a-zA-Z]+', '', s)
    s = re.sub(r'\s+', ' ', s)
    return s.strip()


# --------------------------------------------------------------------------
# docx 构建
# --------------------------------------------------------------------------

def _set_cjk(run, font='等线', size=None, bold=None, italic=None, mono=False):
    name = 'Consolas' if mono else font
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体' if not mono else name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


_INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\$[^$]+\$)')


def add_rich(par, text):
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            r = par.add_run(piece[2:-2])
            _set_cjk(r, bold=True)
        elif piece.startswith('`') and piece.endswith('`'):
            r = par.add_run(piece[1:-1])
            _set_cjk(r, mono=True, size=9.5)
        elif piece.startswith('$') and piece.endswith('$'):
            r = par.add_run(inline_to_text(piece[1:-1]))
            _set_cjk(r, italic=True)
        elif piece.startswith('*') and piece.endswith('*') and len(piece) > 2:
            r = par.add_run(piece[1:-1])
            _set_cjk(r, italic=True)
        else:
            r = par.add_run(piece)
            _set_cjk(r)


def convert(md_path: Path, out_path: Path, base_dir: Path | None = None):
    base_dir = base_dir or md_path.parent
    math_dir = out_path.parent / '_eq'
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = '等线'
    style.font.size = Pt(10.5)
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    lines = md_path.read_text(encoding='utf-8').splitlines()
    i = 0
    fig_no = 0
    table_no = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # 分隔线
        if re.fullmatch(r'-{3,}|\*{3,}|_{3,}', stripped):
            i += 1
            continue

        # 块级公式
        if stripped.startswith('$$'):
            buf = [stripped[2:]]
            if not stripped.endswith('$$') or len(stripped) <= 3:
                i += 1
                while i < len(lines) and '$$' not in lines[i]:
                    buf.append(lines[i])
                    i += 1
                if i < len(lines):
                    buf.append(lines[i].replace('$$', ''))
            else:
                buf = [stripped[2:-2]]
            tex = '\n'.join(buf).strip()
            png = render_math(tex, math_dir)
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if png:
                run = par.add_run()
                run.add_picture(str(png), height=Pt(max(14, 13 + 4 * tex.count('frac'))))
            else:
                r = par.add_run(inline_to_text(tex))
                _set_cjk(r, mono=True, size=10)
            i += 1
            continue

        # 代码块
        if stripped.startswith('```'):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                buf.append(lines[i])
                i += 1
            par = doc.add_paragraph()
            r = par.add_run('\n'.join(buf))
            _set_cjk(r, mono=True, size=9)
            par.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue

        # 图片
        m = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if m:
            caption, src = m.group(1), m.group(2)
            img = (base_dir / src).resolve()
            if img.is_file():
                fig_no += 1
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.add_run().add_picture(str(img), width=Inches(5.9))
                cap = doc.add_paragraph()
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = cap.add_run(caption if caption else f'图 {fig_no}')
                _set_cjk(r, size=9)
                r.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # 表格
        if stripped.startswith('|') and i + 1 < len(lines) and \
                re.fullmatch(r'\|[\s:|-]+\|', lines[i + 1].strip()):
            header = [c.strip() for c in stripped.strip('|').split('|')]
            i += 2
            body = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                body.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            table_no += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = 'Light Grid Accent 1'
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for c, text in zip(t.rows[0].cells, header):
                c.text = ''
                add_rich(c.paragraphs[0], text)
                for run in c.paragraphs[0].runs:
                    run.bold = True
                    run.font.size = Pt(9)
            for row in body:
                cells = t.add_row().cells
                for c, text in zip(cells, row + [''] * (len(header) - len(row))):
                    c.text = ''
                    add_rich(c.paragraphs[0], text)
                    for run in c.paragraphs[0].runs:
                        run.font.size = Pt(9)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'(#{1,6})\s+(.*)', stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            h = doc.add_heading(level=min(level, 4))
            h.alignment = (WD_ALIGN_PARAGRAPH.CENTER if level == 1
                           else WD_ALIGN_PARAGRAPH.LEFT)
            add_rich(h, text)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
                _set_cjk(run, bold=True,
                         size={1: 18, 2: 14, 3: 12, 4: 11}.get(level, 11))
            i += 1
            continue

        # 引用
        if stripped.startswith('>'):
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.3)
            add_rich(par, stripped.lstrip('> ').strip())
            for run in par.runs:
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                run.font.size = Pt(9.5)
            i += 1
            continue

        # 列表
        m = re.match(r'^(\s*)([-*+]|\d+\.)\s+(.*)', line)
        if m:
            indent, marker, text = m.group(1), m.group(2), m.group(3)
            style_name = ('List Number' if marker[0].isdigit() else 'List Bullet')
            par = doc.add_paragraph(style=style_name)
            par.paragraph_format.left_indent = Inches(0.3 + 0.25 * (len(indent) // 2))
            add_rich(par, text)
            i += 1
            continue

        if not stripped:
            i += 1
            continue

        par = doc.add_paragraph()
        par.paragraph_format.first_line_indent = Pt(21)
        par.paragraph_format.space_after = Pt(4)
        add_rich(par, stripped)
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('markdown')
    ap.add_argument('-o', '--output')
    args = ap.parse_args()
    md = Path(args.markdown)
    out = Path(args.output) if args.output else md.with_suffix('.docx')
    print('wrote', convert(md, out))


if __name__ == '__main__':
    main()
